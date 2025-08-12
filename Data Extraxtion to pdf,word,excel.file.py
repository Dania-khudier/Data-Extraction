import os
import pandas as pd
from openpyxl import load_workbook
import fitz
import io
from PIL import Image
from docx import Document
import zipfile

def process_excel_file(filename='contacts.xlsx'):
    
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    
    df = pd.read_excel(filename, engine='openpyxl')
    print(df)

    
    all_sheets = pd.read_excel(filename, sheet_name=None, engine='openpyxl')
    all_data = []

    
    for sheet_name, df in all_sheets.items():
        print(f"Sheet content: {sheet_name}")
        df_str = df.astype(str)  # Convert all data to string
        sheet_data = df_str.values.flatten().tolist()  
        all_data.extend(sheet_data)

        
        for item in sheet_data:
            print(item)

    # Prepare directory for extracted images
    output_dir = "extracted_images"
    os.makedirs(output_dir, exist_ok=True)

    # Load workbook to extract images
    wb = load_workbook(filename)

    img_count = 0
    
    for sheet in wb.worksheets:
        for image in sheet._images:
            img_count += 1
            
            ext = image.path.split('.')[-1] if image.path else "png"
            img_path = os.path.join(output_dir, f"image_{img_count}.{ext}")
            # Save image data to file
            with open(img_path, "wb") as f:
                f.write(image._data())
            print(f"Extraction images from file {img_path}")

    print(f"Total extracted images from Excel file: {img_count}")

def process_pdf_file(pdf_path="معلومات شركة COXIR الكورية.pdf"):
    
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    # Open PDF file using fitz
    doc = fitz.open(pdf_path)
    full_text = ""

    # Extract text from all pages
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text()
        full_text += text + "\n"

    # Print all extracted text
    print(full_text)

    
    output_folder = "extracted_images"
    os.makedirs(output_folder, exist_ok=True)

    # Extract images from each page
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        image_list = page.get_images(full=True)

        for img_index, img in enumerate(image_list, start=1):
            xref = img[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]
            image_ext = base_image["ext"]

            # Open image bytes and save to file
            image = Image.open(io.BytesIO(image_bytes))
            image_filename = os.path.join(output_folder, f"page{page_num+1}_img{img_index}.{image_ext}")
            image.save(image_filename)
            print(f"تم حفظ الصورة: {image_filename}")

    # Write all extracted text to a file
    with open("output.txt", "w", encoding="utf-8") as f:
        f.write(full_text)

def process_docx_file(docx_path="Product Requirements Document.docx"):
    
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    # Open DOCX document
    doc = Document(docx_path)
    print("Extract Texts From The File")
    
    for para in doc.paragraphs:
        print(para.text)

    # Prepare output directory for images
    output_dir = "extracted_images"
    
    with zipfile.ZipFile(docx_path, "r") as docx_zip:
        
        media_files = [f for f in docx_zip.namelist() if f.startswith('word/media/')]
        os.makedirs(output_dir, exist_ok=True)
        # Extract all media files to the output directory
        for file in media_files:
            docx_zip.extract(file, output_dir)
            print(f"images extracted to file {output_dir}")


process_excel_file()
process_pdf_file()
process_docx_file()
